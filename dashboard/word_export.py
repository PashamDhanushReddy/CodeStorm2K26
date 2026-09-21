import requests
from io import BytesIO
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

from PIL import Image

def get_logo():
    try:
        r = requests.get('https://www.nrcmec.org/Student/images/NRCM-Logo.png', timeout=5)
        if r.status_code == 200:
            img = Image.open(BytesIO(r.content))
            # Crop the tree icon from the left side (it's 320x132, we just take 132x132)
            cropped = img.crop((0, 0, 132, 132))
            
            buf = BytesIO()
            cropped.save(buf, format='PNG')
            buf.seek(0)
            return buf
    except:
        pass
    return None

def set_cell_bg_color(cell, color_hex):
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color_hex))
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_chart(codes, counts):
    fig, ax = plt.subplots(figsize=(8, 4))
    bars = ax.bar(codes, counts, color='#5870f0', width=0.8, edgecolor='white')
    
    for bar in bars:
        height = bar.get_height()
        ax.annotate(f'{height}',
                    xy=(bar.get_x() + bar.get_width() / 2, height),
                    xytext=(0, 3), 
                    textcoords="offset points",
                    ha='center', va='bottom', fontweight='bold', fontsize=10)
    
    ax.yaxis.grid(True, linestyle='-', which='major', color='lightgrey', alpha=0.7)
    ax.xaxis.grid(False)
    ax.set_axisbelow(True)
    
    ax.set_ylabel('Number of Teams', color='gray')
    ax.set_xlabel('College Code', color='gray')
    
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('gray')
    ax.spines['bottom'].set_color('gray')
    ax.tick_params(colors='gray')
    
    plt.tight_layout()
    buf = BytesIO()
    plt.savefig(buf, format='png', dpi=150, bbox_inches='tight')
    plt.close(fig)
    buf.seek(0)
    return buf

def create_word_report(total_teams, total_colleges, peak_day, col_wise_data):
    doc = Document()
    
    # Set narrow margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # 1. Header (Logo + Text)
    table = doc.add_table(rows=1, cols=2)
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(6.0)
    
    logo_cell = table.cell(0, 0)
    text_cell = table.cell(0, 1)
    
    logo_buf = get_logo()
    if logo_buf:
        para = logo_cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        run.add_picture(logo_buf, width=Inches(1.2))
        
    para = text_cell.paragraphs[0]
    run = para.add_run("NARSIMHA REDDY\nENGINEERING COLLEGE\n")
    run.font.name = 'Arial'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(79, 38, 131)
    
    run2 = para.add_run("An Autonomous Institution | Affiliated to JNTUH | Approved by AICTE\nAccredited by NBA & NAAC with 'A' Grade")
    run2.font.name = 'Arial'
    run2.font.size = Pt(8)
    run2.font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_paragraph()
    
    # CODESTORM 2K26 Title
    p = doc.add_paragraph()
    r = p.add_run("NRCM'S ")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 255)
    
    r2 = p.add_run("CODE")
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 0, 0)
    
    r3 = p.add_run("STORM ")
    r3.font.size = Pt(20)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(79, 102, 238)
    
    r4 = p.add_run("2K26")
    r4.font.size = Pt(20)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(255, 0, 0)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("A 36-HOUR NATIONAL LEVEL HACKATHON • IDEAS TODAY, IMPACT TOMORROW.")
    r_sub.font.size = Pt(9)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)
    
    # Banner
    p_banner = doc.add_paragraph()
    r_ban = p_banner.add_run("TEAM REGISTRATION REPORT")
    r_ban.font.size = Pt(14)
    r_ban.font.bold = True
    
    doc.add_paragraph("A clear snapshot of team registrations by college code and daily registration activity.")
    
    # Stats Table
    stats_table = doc.add_table(rows=2, cols=3)
    stats_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in stats_table.rows:
        for cell in row.cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    h1 = stats_table.cell(0, 0).paragraphs[0].add_run("TOTAL TEAMS")
    h2 = stats_table.cell(0, 1).paragraphs[0].add_run("TOTAL COLLEGES")
    h3 = stats_table.cell(0, 2).paragraphs[0].add_run("PEAK REGISTRATION DAY")
    for h in [h1, h2, h3]:
        h.font.bold = True
        h.font.size = Pt(9)
        h.font.color.rgb = RGBColor(150, 150, 150)
        
    v1 = stats_table.cell(1, 0).paragraphs[0].add_run(str(total_teams))
    v1.font.color.rgb = RGBColor(0, 112, 192)
    v2 = stats_table.cell(1, 1).paragraphs[0].add_run(str(total_colleges))
    v2.font.color.rgb = RGBColor(112, 48, 160)
    v3 = stats_table.cell(1, 2).paragraphs[0].add_run(str(peak_day))
    v3.font.color.rgb = RGBColor(255, 0, 0)
    
    for v in [v1, v2, v3]:
        v.font.bold = True
        v.font.size = Pt(18)
        
    doc.add_paragraph()
    
    # College-wise Registrations
    p_heading = doc.add_paragraph()
    r_h1 = p_heading.add_run("1. ")
    r_h1.font.color.rgb = RGBColor(68, 114, 196)
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h2 = p_heading.add_run("College-wise Registrations")
    r_h2.font.color.rgb = RGBColor(0, 0, 0)
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    
    p_desc = doc.add_paragraph()
    p_desc.add_run("The graph uses ").font.color.rgb = RGBColor(120, 120, 120)
    r_desc_b = p_desc.add_run("college codes")
    r_desc_b.font.color.rgb = RGBColor(79, 38, 131)
    r_desc_b.font.bold = True
    p_desc.add_run(" so the chart stays compact and easy to compare.").font.color.rgb = RGBColor(120, 120, 120)
    
    p_ctitle = doc.add_paragraph()
    p_ctitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ctitle = p_ctitle.add_run("Teams by College Code")
    r_ctitle.font.bold = True
    
    # Chart Image
    codes = [row[0] for row in col_wise_data]
    counts = [row[2] for row in col_wise_data]
    if len(codes) > 0:
        chart_buf = generate_chart(codes, counts)
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(chart_buf, width=Inches(6.5))
    
    doc.add_paragraph()
    
    # Table
    data_table = doc.add_table(rows=len(col_wise_data)+1, cols=3)
    data_table.style = 'Table Grid'
    
    headers = ['College Code', 'College Name', 'Teams']
    for j, h in enumerate(headers):
        cell = data_table.cell(0, j)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(h)
        set_cell_bg_color(cell, '0B163F')
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
        
    for i, row in enumerate(col_wise_data):
        for j, val in enumerate(row):
            cell = data_table.cell(i+1, j)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(val))
            run.font.size = Pt(9)
            
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
